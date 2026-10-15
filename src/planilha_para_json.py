"""Atualiza src/data/inversao-fluxo.json a partir de uma versão nova da planilha.
Uso: python src/planilha_para_json.py Celesc_inversao_de_fluxo.xlsx  (depois rode src/build.py)
Requer: pip install openpyxl"""
import json, pathlib, sys
import openpyxl

SRC = pathlib.Path(__file__).resolve().parent
FCOLS = ['Alimentador', 'Nome do alimentador', 'Subestação', 'Município principal', 'Norte Catarinense',
         'Transformador da SE', 'UCs ativas', 'Carga mínima ao meio-dia (kW)', 'Mês da carga mínima',
         'Carga média (kW)', 'Qtd. GD', 'Potência GD instalada (kW)', 'Usinas conectadas (kW instalados)',
         'Tipo das usinas', 'Geração média das usinas (kW)', 'Reverso medido no transformador (MWh/ano)',
         'Meses com reverso no transformador']

def rnd(v):
    return round(v, 4) if isinstance(v, float) else v

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
p = wb['Parâmetros']
params = {'inj': p['B2'].value, 't1': p['B3'].value, 't2': p['B4'].value, 't3': p['B5'].value}
a = wb['Alimentadores (estimado)']
col = {c.value: i for i, c in enumerate(a[1])}
feeders = [[rnd(r[col[c]]) for c in FCOLS] for r in a.iter_rows(min_row=2, values_only=True) if r[col['Alimentador']]]
b = wb['Subestações (medido)']
ses = [[rnd(v) for v in r] for r in b.iter_rows(min_row=2, values_only=True) if r[0]]
out = {'params': params, 'fcols': FCOLS, 'feeders': feeders, 'scols': [c.value for c in b[1]], 'ses': ses}
(SRC / 'data/inversao-fluxo.json').write_text(json.dumps(out, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
print(len(feeders), 'alimentadores e', len(ses), 'transformadores de SE')
